# -*- coding: utf-8 -*-
"""Surgical edits on the user's hand-tuned deck: their layout and copy stay, only
the requested slides change."""
import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

A = 'assets'
INK   = RGBColor(0x0C,0x1D,0x28); INK3 = RGBColor(0x1A,0x2F,0x3B)
CREAM = RGBColor(0xF7,0xF1,0xE7); GOLD = RGBColor(0xC8,0x9B,0x45)
GOLDL = RGBColor(0xD9,0xBC,0x7A); GOLDD = RGBColor(0x6E,0x54,0x26)
FF='Inter'

prs = Presentation('src-images/deck-v2-hand-edited.pptx')
S = list(prs.slides)

def txt(sh):
    return " | ".join(p.text for p in sh.text_frame.paragraphs if p.text.strip()) if sh.has_text_frame else ""

def find(slide, needle, nth=0):
    hits=[sh for sh in slide.shapes if sh.has_text_frame and needle in txt(sh)]
    return hits[nth]

def by_id(slide, sid):
    return next(sh for sh in slide.shapes if sh.shape_id==sid)

def set_lines(sh, lines):
    """Replace the text, keeping each paragraph's run formatting."""
    tf = sh.text_frame
    tmpl = next((p for p in tf.paragraphs if p.runs), None)
    if tmpl is None: return
    while len(tf.paragraphs) < len(lines):
        tf._txBody.append(copy.deepcopy(tmpl._p))
    while len(tf.paragraphs) > len(lines):
        tf._txBody.remove(tf.paragraphs[-1]._p)
    for p, line in zip(tf.paragraphs, lines):
        if not p.runs:
            p._p.append(copy.deepcopy(tmpl.runs[0]._r))
        p.runs[0].text = line
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)

def recolor(sh, rgb):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = rgb

def drop(sh):
    sh._element.getparent().remove(sh._element)

def swap_icon(slide, pic, name, size=None):
    l,t,w,h = pic.left, pic.top, pic.width, pic.height
    if size: w=h=Inches(size); l=pic.left+(pic.width-w)//2; t=pic.top+(pic.height-h)//2
    drop(pic)
    return slide.shapes.add_picture(os.path.join(A,f'ic-{name}.png'), l,t,w,h)

# ───────────────────────── 04 · три пути в деньгах ─────────────────────────
s = S[3]
set_lines(find(s,'ВЫВЕЗТИ ВСЁ'), ['ВЫБРОСИТЬ ВСЁ'])
set_lines(find(s,'Демонтаж, машины'), ['Вывоз и утилизация','за ваш счёт'])
set_lines(find(s,'ИТ — одним лотом'), ['ИТ — за полцены.','Мебель — на вывоз'])
set_lines(find(s,'Продано 420 000 ₽'), ['420 000 ₽ − услуга 172 000 ₽','(25 000 ₽ + 35%)'])
set_lines(find(s,'Пример по референсному'), ['Пример по референсному объекту.'])

# ───────────────────────── 05 · механика ─────────────────────────
s = S[4]
steps = [
    ('ФОТО',        ['СООБЩАЕТЕ'],    'и дата освобождения', ['от чего нужно избавиться'], 'messagesquare'),
    ('ОПИСЬ',       ['СЧИТАЕМ'],      'Придем и отснимем',   ['сколько выручим'],          'barchart2'),
    ('ТЕНДЕР',      ['ПОДПИСЫВАЕМ'],  'ИТ, мебель',          ['договор'],                  'edit3'),
    ('ВЫВОЗ',       ['ПРОДАЁМ'],      'демонтаж, пропуска',  ['ваши вещи'],                'tag'),
    ('ОТЧЁТ',       ['ОТЧИТЫВАЕМСЯ'], 'кому продано',        ['и делим прибыль'],          'piechart'),
]
pics = [sh for sh in s.shapes if sh.shape_type==13]
for (t_old,t_new,c_old,c_new,icon), pic in zip(steps, pics):
    set_lines(find(s,t_old), t_new)
    set_lines(find(s,c_old), c_new)
    swap_icon(s, pic, icon)

# ───────────────────────── 06 · стоимость услуги ─────────────────────────
s = S[5]
set_lines(find(s,'ТАРИФЫ'), ['ДВА ФОРМАТА'])
set_lines(find(s,'ДВА ФОРМАТА',1), ['СТОИМОСТЬ УСЛУГИ'])
sub = find(s,'Выбираете по одному признаку'); set_lines(sub,['Выбираете по задаче']); recolor(sub, CREAM)
e1 = find(s,'ФОРМАТ A');              set_lines(e1,['ДЛЯ МАКСИМАЛЬНОЙ ВЫРУЧКИ']); recolor(e1, GOLD)
e2 = find(s,'ФОРМАТ B · РЕКОМЕНДУЕМ');set_lines(e2,['ДЛЯ БЫСТРОГО ОСВОБОЖДЕНИЯ']); recolor(e2, GOLD)
set_lines(find(s,'Имущество остаётся'), ['Вещи остаются в офисе до продажи'])
set_lines(find(s,'Оценка, лоты'),       ['Продаём дольше — выручаем больше'])
set_lines(find(s,'Помещение свободно'), ['Помещение свободно к вашей дате'])
set_lines(find(s,'Демонтаж, вывоз'),    ['Продаём дальше без вашего участия'])
for sid in (17,18,30,31):               # третьи буллеты в обеих карточках
    drop(by_id(s,sid))
for sid in (10,23):                     # обе цены одним цветом — ни один формат не выделен
    recolor(by_id(s,sid), GOLDL)
for sid in (11,24):                     # «от суммы продаж» — контрастнее
    recolor(by_id(s,sid), CREAM)
for sid in (6,19):                      # одинаковые карточки: одна заливка, одна рамка
    c = by_id(s,sid)
    c.fill.solid(); c.fill.fore_color.rgb = INK3
    c.line.color.rgb = GOLDD; c.line.width = Pt(1)

# ───────────────────────── 07 · кресло целиком ─────────────────────────
s = S[6]
img = by_id(s,12); frame = by_id(s,13)
new_h = Emu(int(Inches(4.54)/2.043))
drop(img); drop(frame)
s.shapes.add_picture(os.path.join(A,'ph-P8-chair.jpg'), Inches(7.52), Inches(1.66), Inches(4.54), new_h)
by_id(s,14).top = Inches(4.02)
for k,sid in enumerate((15,16)): by_id(s,sid).top = Inches(4.48)
by_id(s,17).top = Inches(4.81)
for sid in (18,19): by_id(s,sid).top = Inches(4.88)
by_id(s,20).top = Inches(5.21)
for sid in (21,22): by_id(s,sid).top = Inches(5.28)
by_id(s,23).top = Inches(5.61)
for sid in (24,25): by_id(s,sid).top = Inches(5.68)
by_id(s,26).top = Inches(6.01)
by_id(s,27).top = Inches(6.14); by_id(s,28).top = Inches(6.21)

# ───────────────────────── 10 · расчёт бесплатно ─────────────────────────
s = S[9]
set_lines(find(s,'НАЧАЛО'), ['ПЕРВЫЙ ШАГ'])
set_lines(find(s,'СТАРТ В 2 ШАГА'), ['РАСЧЁТ — БЕСПЛАТНО'])
set_lines(find(s,'ОТПРАВЬТЕ') if any('ОТПРАВЬТЕ' in txt(x) for x in s.shapes) else find(s,'ПРИШЛИТЕ'), ['ОТПРАВЬТЕ'])
set_lines(find(s,'3–5 минут видео'), ['фото или просто названия того,','что нужно продать'])
set_lines(find(s,'маршрут имущества'), ['расчёт и маршрут имущества'])
set_lines(find(s,'БЕСПЛАТНАЯ'), ['ОЦЕНКА','ПО ФОТО','ИЛИ СПИСКУ'])
swap_icon(s, by_id(s,14), 'camera2')
drop(find(s,'Отдельная большая инвентаризация'))

# ───────────────────────── 11 · только неочевидные вопросы ─────────────────────────
s = S[10]
for sid in (20,21,22,23,24,25,26,27): drop(by_id(s,sid))
qa = [
    (4,5,6,   '01','Что с непроданным?',                'Финальный маршрут согласуем заранее: передача, вторсырьё или утилизация.'),
    (8,9,10,  '02','А данные на технике?',              'Учётки отвязываем, накопители стираем или изымаем до передачи.'),
    (12,13,14,'03','Имущество на балансе или в лизинге?','Сначала разложим, чем компания вправе распоряжаться, и только потом публикуем.'),
    (16,17,18,'04','Как оформляется оплата?',           'Договор, счёт, безналичный расчёт и закрывающие документы.'),
]
for n_id,q_id,a_id,n,q,a in qa:
    set_lines(by_id(s,n_id),[n]); set_lines(by_id(s,q_id),[q]); set_lines(by_id(s,a_id),[a])
rows = {0:(2.54,2.50,3.22,3.98), 1:(4.46,4.42,5.14,5.90)}
for i,(n_id,q_id,a_id,_,_,_) in enumerate(qa):
    ny,qy,ay,dy = rows[i//2]
    by_id(s,n_id).top=Inches(ny); by_id(s,q_id).top=Inches(qy); by_id(s,a_id).top=Inches(ay)
for i,sid in enumerate((7,11,15,19)):
    by_id(s,sid).top = Inches(rows[i//2][3])

# ───────────────────────── 12 · контакты и три QR ─────────────────────────
s = S[11]
set_lines(find(s,'+7 965'), ['+7 917 896-84-83'])
set_lines(find(s,'ODRybakov'), ['proday_za_menya@mail.ru'])
for sid in (12,13,14,15,16,17,18,19,20,21,22,23): drop(by_id(s,sid))

cw, gap, x0, cy, ch = 1.75, 0.14, 6.95, 2.68, 2.92
for i,(lab,qr,icon) in enumerate([('TELEGRAM','qr-telegram.png','send'),
                                  ('MAX','qr-max.png','messagesquare'),
                                  ('WHATSAPP','qr-whatsapp.png','messagecircle')]):
    x = x0 + i*(cw+gap)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(cy), Inches(cw), Inches(ch))
    card.adjustments[0] = 0.06
    card.fill.solid(); card.fill.fore_color.rgb = INK
    card.line.color.rgb = GOLDD; card.line.width = Pt(1)
    card.shadow.inherit = False
    card.text_frame.text = ''
    s.shapes.add_picture(os.path.join(A,f'ic-{icon}.png'), Inches(x+(cw-0.34)/2), Inches(cy+0.22), Inches(0.34), Inches(0.34))
    plate = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x+(cw-1.42)/2), Inches(cy+0.74), Inches(1.42), Inches(1.42))
    plate.adjustments[0] = 0.05
    plate.fill.solid(); plate.fill.fore_color.rgb = CREAM
    plate.line.fill.background(); plate.shadow.inherit = False
    plate.text_frame.text = ''
    s.shapes.add_picture(os.path.join(A,qr), Inches(x+(cw-1.36)/2), Inches(cy+0.77), Inches(1.36), Inches(1.36))
    tb = s.shapes.add_textbox(Inches(x), Inches(cy+2.38), Inches(cw), Inches(0.28))
    tf = tb.text_frame; tf.word_wrap=False; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = lab
    r.font.size=Pt(11); r.font.bold=True; r.font.name=FF; r.font.color.rgb=GOLDL

# ───────────────────────── 08 · подрезать хвосты текста ─────────────────────────
s = S[7]
set_lines(find(s,'согласованное имущество'), ['согласованное имущество — к вашей дате'])
set_lines(find(s,'дальше работаем с покупателями'), ['работаем с покупателями без вас'])
set_lines(find(s,'финальный маршрут согласован'), ['передача, вторсырьё или утилизация'])

# ───────────────────────── 03 · контраст подписей ─────────────────────────
s = S[2]
for needle in ('После решения', 'Заплатить за вывоз', 'Опись, фото, объявления', 'Помещение пустое к дате'):
    recolor(find(s, needle), CREAM)

# ───────────────────────── вертикальный баланс ─────────────────────────
for sid in range(4,30):                       # 05: весь ряд этапов ниже
    try: sh=by_id(S[4],sid)
    except StopIteration: continue
    sh.top = sh.top + Inches(0.40)
for sh in S[4].shapes:                         # иконки добавлены заново — сдвинуть их тоже
    if sh.shape_type==13 and Emu(sh.top).inches < 3.0:
        sh.top = sh.top + Inches(0.40)
for sid in (5,6,7,8,9,10,11,12,13,15):         # 10: шаги и карточка ниже
    by_id(S[9],sid).top = by_id(S[9],sid).top + Inches(0.42)
for sh in S[9].shapes:
    if sh.shape_type==13 and Emu(sh.top).inches < 4.0:
        sh.top = sh.top + Inches(0.42)

prs.save('Продай_за_меня_ОФИСЫ_клиентская.pptx')
print('saved')
